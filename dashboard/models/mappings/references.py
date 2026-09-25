"""Read-only Martini fragment reference index for mapping suggestions."""
from __future__ import annotations

from pathlib import Path
import hashlib
import json
import re

from .identity import canonical_sha256


TYPE_RE = re.compile(r"(?<![A-Za-z0-9])([STR]?[PQNCX][1-6][a-z]?)(?![A-Za-z0-9])")


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _size(bead_type: str) -> str:
    return bead_type[0] if bead_type[:1] in {"S", "T"} else "R"


def _record(fragment_id, bead_type, rank, path, locator, *, name=None, notes=None):
    return {
        "fragment_id": str(fragment_id), "type": bead_type, "size": _size(bead_type),
        "source_rank": rank, "name": name or str(fragment_id),
        "notes": notes or "", "references": [{"path": str(path), "sha256": _sha(path), "locator": locator}],
    }


def _parse_itp(path: Path, rank: int, source_name: str):
    records = []
    section = ""; molecule = None; seen = set()
    for line_number, raw in enumerate(path.read_text(errors="replace").splitlines(), 1):
        line = raw.split(";", 1)[0].strip()
        if line.startswith("["):
            section = line.strip("[] ").lower(); continue
        if not line or line.startswith("#"):
            continue
        fields = line.split()
        if section == "moleculetype":
            molecule = fields[0]
        elif section == "atoms" and len(fields) >= 2:
            bead_type = fields[1]
            if not TYPE_RE.fullmatch(bead_type) or (molecule, bead_type) in seen:
                continue
            seen.add((molecule, bead_type))
            records.append(_record(f"{source_name}:{molecule}:{bead_type}", bead_type, rank, path,
                                   f"line:{line_number}", name=molecule,
                                   notes="Martini topology bead type; atom attachment requires a verified AA mapping."))
    return records


def _parse_sugars(path: Path):
    return _parse_itp(path, 1, "sugars")


def _parse_bible(path: Path):
    records = []
    try:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
    except (ImportError, OSError, ValueError) as exc:
        return records, [f"Bible workbook unavailable: {exc}"]
    try:
        sheet = book["All"]
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            name, tag, smiles, cgsmiles = row[:4]
            if not tag or not cgsmiles:
                continue
            for bead_type in sorted(set(TYPE_RE.findall(str(cgsmiles)))):
                records.append(_record(f"bible:{tag}:{bead_type}", bead_type, 3, path,
                                       f"All!row:{row_number}", name=str(name or tag),
                                       notes=f"M3-Bible CGSmiles fragment; SMILES={smiles or 'unrecorded'}"))
    finally:
        book.close()
    return records, []


def load_index(project_root: Path) -> dict:
    """Index known reference files without copying or mutating them."""
    root = Path(project_root); records = []; sources = []; diagnostics = []
    sugar_paths = sorted((root / "references").glob("martini_*_sugars_*.itp"))
    for path in sugar_paths:
        if path.is_symlink() or not path.is_file():
            continue
        sources.append({"path": str(path.relative_to(root)), "sha256": _sha(path), "rank": 1})
        records.extend(_parse_sugars(path))
    small_root = root / "references" / "M3-Small-Molecules" / "models"
    small_itp = small_root / "martini_v3.0.0_small_molecules_v1.itp"
    if small_itp.is_file() and not small_itp.is_symlink():
        sources.append({"path": str(small_itp.relative_to(root)), "sha256": _sha(small_itp), "rank": 2})
        records.extend(_parse_itp(small_itp, 2, "small-molecules"))
    bible = root / "references" / "M3-Bible" / "src" / "MartiniCGSmilesDB_v6.xlsx"
    if bible.is_file() and not bible.is_symlink():
        sources.append({"path": str(bible.relative_to(root)), "sha256": _sha(bible), "rank": 3})
        bible_records, bible_diagnostics = _parse_bible(bible)
        records.extend(bible_records); diagnostics.extend(bible_diagnostics)
    records.sort(key=lambda item: (item["source_rank"], item["fragment_id"]))
    index_payload = {"sources": sources, "records": records, "diagnostics": diagnostics}
    return index_payload | {"index_sha256": canonical_sha256(index_payload)}
